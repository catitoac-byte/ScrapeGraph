"""
Varredura de lojas de varejo no Google Maps (projeto C&A scrape).

Para cada marca e cidade, descobre todas as lojas na busca do Maps e abre
cada ficha para extrair o que o Google mostra: endereco, telefone, site,
horario de funcionamento, horarios de pico por hora, nota, distribuicao de
estrelas, topicos citados, atributos da aba Sobre e as avaliacoes em texto.

Os horarios de pico nao sao imagem. Cada barra do grafico tem aria-label
"Movimento as HH:00: NN%." e o container .C7xf8b tem sete blocos na ordem
domingo..sabado, com um aviso "Fechada aos ..." no lugar das barras quando a
loja nao abre. Tudo sai como numero.

Armadilhas encontradas na sondagem:
- Chromium headless nao recebe o bloco de horarios de pico. Com janela
  visivel ele aparece. Por isso o padrao e headless=False.
- A busca comum do google.com cai em /sorry (CAPTCHA). O Maps nao.
- page.goto com wait_until="load" estoura o tempo no Maps. Usar
  domcontentloaded e esperar os seletores.
"""

from __future__ import annotations

import csv
import html
import json
import random
import re
import time
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import quote_plus, unquote

from playwright.sync_api import Page, sync_playwright

RAIZ = Path(__file__).resolve().parents[2]
SAIDA = RAIZ / "dados" / "saida" / "lojas_google"

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36"
)
DIAS = ["domingo", "segunda-feira", "terca-feira", "quarta-feira",
        "quinta-feira", "sexta-feira", "sabado"]

# Rola todo elemento rolavel do painel lateral. O container muda de classe
# entre versoes do Maps, entao nao dependemos do nome dele.
JS_ROLAR = """(dy)=>{document.querySelectorAll('div[role="main"] *').forEach(e=>{
  if(e.scrollHeight>e.clientHeight+50 && getComputedStyle(e).overflowY!='visible') e.scrollBy(0,dy)})}"""

_ESPACOS = re.compile(r"\s+")


def _limpar(texto: str | None) -> str:
    return _ESPACOS.sub(" ", html.unescape(texto)).strip() if texto else ""


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto.lower())
                   if unicodedata.category(c) != "Mn")


def _num(texto: str | None) -> float | None:
    if not texto:
        return None
    m = re.search(r"\d[\d.]*(?:,\d+)?", texto)
    if not m:
        return None
    return float(m.group().replace(".", "").replace(",", "."))


@dataclass
class Avaliacao:
    id: str
    autor: str
    autor_info: str
    estrelas: int | None
    data_relativa: str
    data_estimada: str
    texto: str
    curtidas: int | None
    resposta_proprietario: str
    fotos: int


@dataclass
class Loja:
    marca: str
    nome: str
    url: str
    place_id: str = ""
    latitude: float | None = None
    longitude: float | None = None
    categoria: str = ""
    nota: float | None = None
    total_avaliacoes: int | None = None
    endereco: str = ""
    cidade: str = ""
    localizada_em: str = ""
    telefone: str = ""
    site: str = ""
    plus_code: str = ""
    status_agora: str = ""
    horario_funcionamento: dict[str, str] = field(default_factory=dict)
    horarios_pico: dict[str, dict[str, int]] = field(default_factory=dict)
    movimento_agora: str = ""
    distribuicao_estrelas: dict[str, int] = field(default_factory=dict)
    topicos_avaliacoes: dict[str, int] = field(default_factory=dict)
    destaques: list[str] = field(default_factory=list)
    atributos: dict[str, list[str]] = field(default_factory=dict)
    postagem_proprietario: str = ""
    total_fotos: str = ""
    foto_capa: str = ""
    avaliacoes: list[Avaliacao] = field(default_factory=list)
    ordenacao: str = "relevancia"
    coletado_em: str = ""
    erro: str = ""


# ---------------------------------------------------------------- navegador

def abrir_navegador(pw, headless: bool = False):
    browser = pw.chromium.launch(headless=headless)
    ctx = browser.new_context(
        locale="pt-BR", timezone_id="America/Sao_Paulo", user_agent=UA,
        viewport={"width": 1400, "height": 1000},
    )
    return browser, ctx.new_page()


def _ir(page: Page, url: str, espera: float = 4.0) -> None:
    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    if "consent.google" in page.url:
        # Opcao mais restritiva de privacidade
        page.get_by_role("button", name=re.compile("Rejeitar tudo|Reject all")).first.click()
        page.wait_for_load_state("domcontentloaded")
    page.wait_for_timeout(int(espera * 1000))
    if "/sorry/" in page.url:
        raise RuntimeError("Google pediu CAPTCHA. Pare, espere e rode de novo mais devagar.")


# ---------------------------------------------------------------- descoberta

# Fichas que carregam o nome da marca mas nao sao loja (posto, drogaria...)
EXCLUIR_PADRAO = r"posto|combustivel|drogaria|farmacia|caixa eletronico|centro de distribuicao|escritorio|\bcd\b"


def _casa_marca(nome: str, padrao: str, excluir: str = EXCLUIR_PADRAO) -> bool:
    n = _sem_acento(nome)
    return re.search(padrao, n) is not None and not (excluir and re.search(excluir, n))


def descobrir(page: Page, marca: str, padrao: str, consultas: list[str],
              delay: float = 2.0) -> list[str]:
    """Devolve as URLs das fichas de todas as consultas, sem duplicar."""
    achados: dict[str, str] = {}
    for q in consultas:
        _ir(page, f"https://www.google.com/maps/search/{quote_plus(q)}?hl=pt-BR&gl=br")
        if "/maps/place/" in page.url:
            # Resultado unico: o Maps abre a ficha direto
            nome = _limpar(page.locator("h1").first.inner_text())
            if _casa_marca(nome, padrao):
                achados.setdefault(_chave(page.url), page.url)
            continue
        feed = page.locator('div[role="feed"]')
        if not feed.count():
            continue
        parado = 0
        anterior = -1
        while parado < 3:
            feed.evaluate("e=>e.scrollBy(0,5000)")
            page.wait_for_timeout(1500)
            n = page.locator("a.hfpxzc").count()
            parado = parado + 1 if n == anterior else 0
            anterior = n
            if page.get_by_text(re.compile("chegou ao final da lista")).count():
                break
        for nome, href in page.eval_on_selector_all(
            "a.hfpxzc", 'els=>els.map(e=>[e.getAttribute("aria-label")||"", e.href])'
        ):
            if _casa_marca(nome, padrao):
                achados.setdefault(_chave(href), href)
        print(f"      '{q}': {len(achados)} acumuladas")
        time.sleep(delay + random.random())
    return list(achados.values())


def _chave(url: str) -> str:
    m = re.search(r"!1s(0x[0-9a-f]+:0x[0-9a-f]+)", url)
    return m.group(1) if m else url.split("?")[0]


# ---------------------------------------------------------------- mercado

@dataclass
class Negocio:
    """Um resultado da busca do Maps, lido do cartao da lista, sem abrir a ficha."""
    nome: str
    url: str
    place_id: str
    categoria: str = ""
    nota: float | None = None
    total_avaliacoes: int | None = None
    endereco: str = ""
    telefone: str = ""
    latitude: float | None = None
    longitude: float | None = None
    consultas: list[str] = field(default_factory=list)


_NOTA_CARTAO = re.compile(r"^(\d,\d)\(([\d.]+)\)$")
_FONE = re.compile(r"\(\d{2}\) ?\d{4,5}-\d{4}")


def _ler_cartao(nome: str, href: str, texto: str) -> Negocio:
    n = Negocio(nome=_limpar(nome), url=href.split("?")[0], place_id=_chave(href))
    m = re.search(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", unquote(href))
    if m:
        n.latitude, n.longitude = float(m.group(1)), float(m.group(2))
    linhas = [_limpar(l) for l in texto.splitlines() if _limpar(l)]
    for l in linhas:
        mn = _NOTA_CARTAO.match(l)
        if mn and n.nota is None:
            n.nota = _num(mn.group(1))
            n.total_avaliacoes = int(mn.group(2).replace(".", ""))
        elif "·" in l and not n.categoria and not re.match(r"(Aberto|Fechado|Abre|Fecha)", l):
            partes = [p.strip() for p in l.split("·")]
            n.categoria = partes[0]
            n.endereco = partes[-1] if len(partes) > 1 else ""
        fone = _FONE.search(l)
        if fone and not n.telefone:
            n.telefone = fone.group()
    if n.nota is None and re.search(r"Nenhuma avaliação", texto):
        n.total_avaliacoes = 0
    return n


def _colher(page: Page, q: str, achados: dict[str, Negocio]) -> int:
    """Le os cartoes visiveis agora e junta em `achados`. Devolve quantos sao novos."""
    cartoes = page.evaluate("""()=>[...document.querySelectorAll('div[role="feed"] a.hfpxzc')].map(a=>{
      const c=a.closest('div.Nv2PK')||a.parentElement;
      return [a.getAttribute('aria-label')||'', a.href, c?c.innerText:'']})""")
    novos = 0
    for nome, href, texto in cartoes:
        neg = _ler_cartao(nome, href, texto)
        atual = achados.get(neg.place_id)
        if atual is None:
            achados[neg.place_id] = atual = neg
            novos += 1
        else:
            # Cartao esvaziado numa leitura anterior: completa o que faltou
            for campo in ("nota", "total_avaliacoes", "categoria", "endereco", "telefone"):
                if getattr(atual, campo) in (None, "") and getattr(neg, campo) not in (None, ""):
                    setattr(atual, campo, getattr(neg, campo))
        if q not in atual.consultas:
            atual.consultas.append(q)
    return novos


def mapear(page: Page, consultas: list[str], delay: float = 3.0,
           achados: dict[str, Negocio] | None = None) -> dict[str, Negocio]:
    """Varre a lista de resultados de cada consulta e guarda todos os negocios,
    sem filtro de marca. Serve para ver o mercado antes de escolher quem entra
    no painel. So le os cartoes: nao abre fichas nem avaliacoes.

    O Maps esvazia os cartoes do topo quando a lista rola ate o fim (sobra o
    nome, some a nota). Por isso os cartoes sao lidos a cada rolagem."""
    achados = {} if achados is None else achados
    for q in consultas:
        _ir(page, f"https://www.google.com/maps/search/{quote_plus(q)}?hl=pt-BR&gl=br")
        feed = page.locator('div[role="feed"]')
        if not feed.count():
            print(f"      '{q}': sem lista")
            continue
        novos = _colher(page, q, achados)
        parado, anterior = 0, -1
        while parado < 3:
            feed.evaluate("e=>e.scrollBy(0,5000)")
            page.wait_for_timeout(1500)
            novos += _colher(page, q, achados)
            n = page.locator("a.hfpxzc").count()
            parado = parado + 1 if n == anterior else 0
            anterior = n
            if page.get_by_text(re.compile("chegou ao final da lista")).count():
                break
        print(f"      '{q}': {anterior} na lista, {novos} novos, {len(achados)} no total")
        time.sleep(delay + random.random() * 2)
    return achados


# ---------------------------------------------------------------- ficha

def _rolar_painel(page: Page, vezes: int = 12) -> None:
    for _ in range(vezes):
        page.evaluate(JS_ROLAR, 700)
        page.wait_for_timeout(450)


def _info_item(page: Page, seletor: str) -> str:
    el = page.locator(f'div[role="main"] {seletor}').first
    if not el.count():
        return ""
    rotulo = el.get_attribute("aria-label") or ""
    return _limpar(rotulo.split(":", 1)[1] if ":" in rotulo else rotulo)


# Faixas de CEP que o Google as vezes rotula com a cidade vizinha.
# Ex.: Riachuelo do Buriti Shopping aparece como "Goiania" com CEP 74916.
FAIXAS_CEP = [(74900000, 74999999, "Aparecida de Goiânia")]


def _corrigir_cidade(cidade: str, endereco: str) -> str:
    m = re.search(r"\b(\d{5})-?(\d{3})\b", endereco)
    if m:
        cep = int(m.group(1) + m.group(2))
        for ini, fim, nome in FAIXAS_CEP:
            if ini <= cep <= fim:
                return nome
    return cidade


def _extrair_visao_geral(page: Page, loja: Loja) -> None:
    main = page.locator('div[role="main"]').first
    main.wait_for(timeout=20000)
    loja.nome = _limpar(main.get_attribute("aria-label")) or loja.nome
    loja.place_id = _chave(page.url)
    m = re.search(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", unquote(page.url))
    if m:
        loja.latitude, loja.longitude = float(m.group(1)), float(m.group(2))

    cat = page.locator('div[role="main"] button[jsaction*="category"]').first
    loja.categoria = _limpar(cat.inner_text()) if cat.count() else ""

    dados = page.evaluate("""()=>{
      const m=document.querySelector('div[role="main"]');
      const lab=s=>{const e=m.querySelector(s);return e?e.getAttribute('aria-label'):''};
      const o={};
      o.nota=lab('span[aria-label$="estrelas "], span[aria-label*="estrelas"]');
      o.total=lab('span[aria-label$="avaliações"]');
      o.horas=[...m.querySelectorAll('button[aria-label*="Copiar horário"]')].map(e=>e.getAttribute('aria-label'));
      o.dist=[...m.querySelectorAll('tr[aria-label*="estrela"]')].map(e=>e.getAttribute('aria-label'));
      o.topicos=[...m.querySelectorAll('button[aria-label*="mencionado em"]')].map(e=>e.getAttribute('aria-label'));
      o.destaques=[...m.querySelectorAll('a[aria-label^="Foto do avaliador que escreveu"]')].map(e=>e.getAttribute('aria-label'));
      const st=m.querySelector('[data-item-id="oh"], .OMl5r, .o0Svhf');
      o.status=st?st.innerText:'';
      const capa=m.querySelector('button[aria-label^="Foto de"] img');
      o.capa=capa?capa.src:'';
      return o}""")
    loja.nota = _num(dados["nota"])
    total = _num(dados["total"])
    loja.total_avaliacoes = int(total) if total is not None else None
    for rot in dados["horas"]:
        partes = [p.strip() for p in rot.split(",")]
        if len(partes) >= 2:
            dia = _sem_acento(partes[0])
            loja.horario_funcionamento[dia] = partes[1]
    for rot in dados["dist"]:
        m = re.match(r"(\d) estrelas?, ([\d.]+)", rot)
        if m:
            loja.distribuicao_estrelas[m.group(1)] = int(m.group(2).replace(".", ""))
    for rot in dados["topicos"]:
        m = re.match(r"(.+), mencionado em ([\d.]+)", rot)
        if m:
            loja.topicos_avaliacoes[m.group(1)] = int(m.group(2).replace(".", ""))
    loja.destaques = [
        _limpar(d.split("escreveu", 1)[1]).strip('"') for d in dados["destaques"]
    ]
    loja.foto_capa = dados["capa"]

    loja.endereco = _info_item(page, 'button[data-item-id="address"]')
    cidades = re.findall(r"(?:^|,)\s*([^,]+?)\s+-\s+[A-Z]{2}\b", loja.endereco)
    loja.cidade = _corrigir_cidade(_limpar(cidades[-1]) if cidades else "", loja.endereco)
    local = page.locator('div[role="main"] button[data-item-id="locatedin"]').first
    loja.localizada_em = _limpar((local.get_attribute("aria-label") or "").replace(":", "")) if local.count() else ""
    loja.telefone = _info_item(page, 'button[data-item-id^="phone:"]')
    loja.site = _info_item(page, 'a[data-item-id="authority"]')
    loja.plus_code = _info_item(page, 'button[data-item-id="oloc"]')

    texto = main.inner_text()
    m = re.search(r"\n((?:Aberto|Fechado|Abre|Fecha)[^\n]*)", texto)
    loja.status_agora = _limpar(m.group(1)) if m else ""
    m = re.search(r"Do proprietário\n(.+?)\n", texto)
    loja.postagem_proprietario = _limpar(m.group(1)) if m else ""
    m = re.search(r"(Mais de [\d.]+ fotos)", texto)
    loja.total_fotos = _limpar(m.group(1)) if m else ""


def _extrair_pico(page: Page, loja: Loja) -> None:
    blocos = page.evaluate("""()=>{
      const c=document.querySelector('div[role="main"] .C7xf8b');
      if(!c) return null;
      const extra=[...document.querySelectorAll('div[role="main"] [aria-label]')]
        .map(e=>e.getAttribute('aria-label')).filter(a=>/atualmente|agora/i.test(a) && /%/.test(a));
      return {dias:[...c.children].map(d=>({
        fechado:(d.innerText||'').trim(),
        barras:[...d.querySelectorAll('[aria-label]')].map(e=>e.getAttribute('aria-label'))})), extra}}""")
    if not blocos:
        return
    for i, bloco in enumerate(blocos["dias"][:7]):
        horas: dict[str, int] = {}
        for rot in bloco["barras"]:
            m = re.search(r"(\d{2}):00\D+?(\d+)%", rot)
            if m:
                horas[m.group(1)] = int(m.group(2))
            if re.search(r"atual", rot, re.I):
                loja.movimento_agora = _limpar(rot)
        loja.horarios_pico[DIAS[i]] = horas if horas else {}
    if blocos["extra"] and not loja.movimento_agora:
        loja.movimento_agora = _limpar(blocos["extra"][0])


def _abrir_aba(page: Page, prefixo: str) -> bool:
    aba = page.locator(f'button[role="tab"][aria-label^="{prefixo}"]').first
    if not aba.count():
        return False
    aba.click()
    page.wait_for_timeout(2500)
    return True


def _extrair_sobre(page: Page, loja: Loja) -> None:
    if not _abrir_aba(page, "Sobre"):
        return
    secoes = page.evaluate("""()=>[...document.querySelectorAll('div[role="main"] .iP2t7d')].map(s=>({
        titulo:(s.querySelector('h2')||{}).innerText||'',
        itens:[...s.querySelectorAll('li span[aria-label]')].map(e=>e.getAttribute('aria-label'))}))""")
    for s in secoes:
        if s["itens"]:
            loja.atributos[_limpar(s["titulo"])] = [_limpar(i) for i in s["itens"]]


_UNIDADES = {"minuto": 1 / 1440, "hora": 1 / 24, "dia": 1, "semana": 7,
             "mês": 30, "mes": 30, "meses": 30, "ano": 365}


def _estimar_data(relativa: str, agora: datetime) -> str:
    """'há 3 meses', 'um ano atrás' -> data aproximada ISO."""
    r = _sem_acento(relativa)
    m = re.search(r"(\d+|um|uma)\s+(minuto|hora|dia|semana|mes|meses|ano)", r)
    if not m:
        return ""
    n = 1 if m.group(1) in ("um", "uma") else int(m.group(1))
    dias = n * _UNIDADES[m.group(2)]
    return (agora - timedelta(days=dias)).date().isoformat()


def _extrair_avaliacoes(page: Page, loja: Loja, maximo: int) -> None:
    if maximo <= 0 or not _abrir_aba(page, "Avaliações"):
        return
    # Ordena por mais recentes, para a amostra refletir o momento atual
    ordenar = page.locator('button[aria-label="Classificar avaliações"], button[aria-label="Ordenar avaliações"]').first
    if ordenar.count():
        ordenar.click()
        page.wait_for_timeout(1200)
        # Existe um "Mais recentes" escondido na pagina: sem filtrar por
        # visivel, o clique ia para ele e a lista seguia por relevancia.
        item = (page.locator('[role="menuitemradio"]').filter(has_text="Mais recentes")
                .filter(visible=True).first)
        if item.count():
            item.click()
            page.wait_for_timeout(2500)
            loja.ordenacao = "mais recentes"
        else:
            loja.ordenacao = "relevancia (menu de ordenacao nao abriu)"

    parado, anterior = 0, -1
    while parado < 8:
        n = page.locator("div.jftiEf[data-review-id]").count()
        if n >= maximo:
            break
        parado = parado + 1 if n == anterior else 0
        anterior = n
        if parado and parado % 2 == 0:
            # A lista as vezes para de carregar. Subir um pouco e descer de novo
            # dispara o carregamento outra vez.
            page.evaluate(JS_ROLAR, -1500)
            page.wait_for_timeout(600)
            page.locator("div.jftiEf[data-review-id]").last.scroll_into_view_if_needed()
        page.evaluate(JS_ROLAR, 4000)
        page.wait_for_timeout(1300 + 400 * parado)

    carregadas = page.locator("div.jftiEf[data-review-id]").count()
    if carregadas < maximo and page.get_by_text("visualização limitada").count():
        # Limite do Google para acesso sem login. Nao contornamos: fica registrado
        # para refazer a loja mais tarde.
        loja.erro = (f"visualizacao limitada do Google: {carregadas} avaliacoes carregadas, "
                     "refazer mais tarde")

    # Expande textos cortados ("Mais") antes de ler
    page.evaluate("""()=>document.querySelectorAll('div[role="main"] button.w8nwRe').forEach(b=>b.click())""")
    page.wait_for_timeout(800)

    brutas = page.evaluate("""()=>[...document.querySelectorAll('div.jftiEf[data-review-id]')].map(r=>{
      const q=s=>r.querySelector(s);
      const est=q('span[role="img"][aria-label*="estrela"]');
      const like=q('button[aria-label*="Gostei"] span.pkWtMe, span.pkWtMe');
      return {id:r.getAttribute('data-review-id'),
        autor:(q('.d4r55')||{}).innerText||'',
        info:(q('.RfnDt')||{}).innerText||'',
        estrelas:est?est.getAttribute('aria-label'):'',
        nota_txt:(q('.fzvQIb')||{}).innerText||'',
        data:(q('.rsqaWe')||q('.xRkPPb')||{}).innerText||'',
        texto:([...r.querySelectorAll('.wiI7pd')].find(e=>!e.closest('.CDe7pd'))||{}).innerText||'',
        curtidas:like?like.innerText:'',
        resposta:(q('.CDe7pd .wiI7pd')||q('.CDe7pd')||{}).innerText||'',
        fotos:r.querySelectorAll('button.Tya61d').length}})""")
    agora = datetime.now()
    vistos = set()
    for b in brutas[:maximo]:
        if b["id"] in vistos:
            continue
        vistos.add(b["id"])
        estrelas = _num(b["estrelas"]) or _num(b["nota_txt"])
        data = _limpar(b["data"])
        loja.avaliacoes.append(Avaliacao(
            id=b["id"],
            autor=_limpar(b["autor"]),
            autor_info=_limpar(b["info"]),
            estrelas=int(estrelas) if estrelas is not None else None,
            data_relativa=data,
            data_estimada=_estimar_data(data, agora),
            texto=_limpar(b["texto"]),
            curtidas=int(_num(b["curtidas"])) if _num(b["curtidas"]) else None,
            resposta_proprietario=_limpar(b["resposta"]),
            fotos=b["fotos"],
        ))


def raspar_loja(page: Page, marca: str, url: str, max_avaliacoes: int,
                cidade_alvo: str = "") -> Loja:
    sep = "&" if "?" in url else "?"
    loja = Loja(marca=marca, nome="", url=url.split("?")[0])
    try:
        _ir(page, f"{url}{sep}hl=pt-BR&gl=br", espera=4.5)
        _extrair_visao_geral(page, loja)
        _rolar_painel(page)
        _extrair_pico(page, loja)
        # Topicos e destaques podem carregar so depois da rolagem
        if not loja.topicos_avaliacoes:
            _extrair_visao_geral(page, loja)
        if cidade_alvo and loja.cidade and _sem_acento(loja.cidade) != cidade_alvo:
            # Loja de outro municipio: nao vale abrir abas nem rolar avaliacoes
            return loja
        _extrair_sobre(page, loja)
        _extrair_avaliacoes(page, loja, max_avaliacoes)
    except Exception as e:  # segue para a proxima loja
        loja.erro = f"{type(e).__name__}: {e}"[:300]
    loja.coletado_em = datetime.now(timezone.utc).isoformat(timespec="seconds")
    return loja


# ---------------------------------------------------------------- orquestracao

class AcessoLimitado(RuntimeError):
    """O Google entrou em visualizacao limitada: parar e tentar mais tarde."""


def _limitada(loja: Loja) -> bool:
    return loja.erro.startswith("visualizacao limitada")


def acesso_liberado(page: Page, url_ficha: str) -> bool:
    """Abre a aba de avaliacoes de uma ficha e confere se o Google limitou."""
    _ir(page, f"{url_ficha}{'&' if '?' in url_ficha else '?'}hl=pt-BR&gl=br", espera=4)
    _abrir_aba(page, "Avaliações")
    return page.get_by_text("visualização limitada").count() == 0


def varrer(marcas: dict[str, dict], cidade: str, max_avaliacoes: int,
           headless: bool = False, delay: float = 2.0,
           so_cidade: bool = True, lojas: list[Loja] | None = None) -> list[Loja]:
    """Preenche `lojas` (lista do chamador) para que uma parada por limite
    ainda deixe o que foi coletado disponivel para exportar."""
    lojas = [] if lojas is None else lojas
    with sync_playwright() as pw:
        browser, page = abrir_navegador(pw, headless=headless)
        try:
            for marca, cfg in marcas.items():
                consultas = [c.format(cidade=cidade) for c in cfg["consultas"]]
                print(f"\n[{marca}] descobrindo lojas")
                urls = descobrir(page, marca, cfg["padrao"], consultas, delay)
                print(f"[{marca}] {len(urls)} fichas para abrir")
                alvo = _sem_acento(cidade.split(",")[0].strip())
                if urls and not acesso_liberado(page, urls[0]):
                    raise AcessoLimitado("Google em visualizacao limitada antes de comecar")
                seguidas = 0
                for i, url in enumerate(urls, 1):
                    loja = raspar_loja(page, marca, url, max_avaliacoes,
                                       cidade_alvo=alvo if so_cidade else "")
                    fora = so_cidade and loja.cidade and _sem_acento(loja.cidade) != alvo
                    marca_status = "FORA DA CIDADE, descartada" if fora else (loja.erro or "ok")
                    print(f"   {i:>2}/{len(urls)} {loja.nome[:28]:<28} {loja.cidade[:20]:<20} "
                          f"pico={len(loja.horarios_pico)}d aval={len(loja.avaliacoes)} {marca_status}")
                    if not fora:
                        lojas.append(loja)
                    seguidas = seguidas + 1 if _limitada(loja) else 0
                    if seguidas >= 2:
                        raise AcessoLimitado(f"duas lojas seguidas limitadas em {marca}")
                    time.sleep(delay + random.random() * 2)
        finally:
            browser.close()
    return lojas


# ---------------------------------------------------------------- exportacao

def _escrever_csv(caminho: Path, linhas: list[dict]) -> None:
    if not linhas:
        return
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()))
        w.writeheader()
        w.writerows(linhas)


def tabelas(lojas: list[Loja]) -> dict[str, list[dict]]:
    t: dict[str, list[dict]] = {k: [] for k in
        ("lojas", "horario_funcionamento", "horarios_pico", "pico_resumo", "pico_marca",
         "avaliacoes", "topicos", "atributos")}
    for lj in lojas:
        chave = {"marca": lj.marca, "loja": lj.nome, "place_id": lj.place_id,
                 "localizada_em": lj.localizada_em}
        media_dist = sum(lj.distribuicao_estrelas.values()) or 0
        t["lojas"].append({
            **chave, "categoria": lj.categoria, "nota": lj.nota,
            "total_avaliacoes": lj.total_avaliacoes,
            **{f"estrelas_{k}": lj.distribuicao_estrelas.get(k) for k in "54321"},
            "pct_negativas_1e2": round(100 * (lj.distribuicao_estrelas.get("1", 0)
                + lj.distribuicao_estrelas.get("2", 0)) / media_dist, 1) if media_dist else None,
            "endereco": lj.endereco, "cidade": lj.cidade, "telefone": lj.telefone,
            "site": lj.site, "plus_code": lj.plus_code, "latitude": lj.latitude,
            "longitude": lj.longitude, "status_na_coleta": lj.status_agora,
            "movimento_agora": lj.movimento_agora, "total_fotos": lj.total_fotos,
            "destaques": " | ".join(lj.destaques),
            "postagem_proprietario": lj.postagem_proprietario,
            "avaliacoes_coletadas": len(lj.avaliacoes), "url": lj.url,
            "foto_capa": lj.foto_capa, "coletado_em": lj.coletado_em, "erro": lj.erro,
        })
        for dia in DIAS:
            hf = lj.horario_funcionamento.get(dia) or lj.horario_funcionamento.get(
                dia.replace("terca", "terça").replace("sabado", "sábado"), "")
            t["horario_funcionamento"].append({**chave, "dia": dia, "horario": hf})
        for dia, horas in lj.horarios_pico.items():
            for hora, pct in sorted(horas.items()):
                t["horarios_pico"].append({**chave, "dia": dia, "hora": int(hora), "movimento_pct": pct})
            abertas = {h: p for h, p in horas.items() if p > 0}
            pico_h = max(abertas, key=abertas.get) if abertas else None
            t["pico_resumo"].append({
                **chave, "dia": dia, "fechado_ou_sem_dado": not abertas,
                "hora_pico": int(pico_h) if pico_h else None,
                "movimento_pico_pct": abertas.get(pico_h) if pico_h else None,
                "horas_mais_cheias": ", ".join(f"{h}h ({abertas[h]}%)" for h in
                    sorted(abertas, key=abertas.get, reverse=True)[:3]),
                "movimento_medio_pct": round(sum(abertas.values()) / len(abertas), 1) if abertas else None,
                "indice_dia": sum(abertas.values()),
            })
        for a in lj.avaliacoes:
            t["avaliacoes"].append({**chave, **asdict(a)})
        for topico, n in lj.topicos_avaliacoes.items():
            t["topicos"].append({**chave, "topico": topico, "mencoes": n})
        for secao, itens in lj.atributos.items():
            for item in itens:
                t["atributos"].append({**chave, "secao": secao, "atributo": item})
    t["pico_marca"] = pico_por_marca(lojas)
    return t


def pico_por_marca(lojas: list[Loja], top: int = 3) -> list[dict]:
    """Curva media por marca, dia e hora (so lojas abertas naquela hora)."""
    soma: dict[tuple[str, str, int], list[int]] = {}
    for lj in lojas:
        for dia, horas in lj.horarios_pico.items():
            for hora, pct in horas.items():
                if pct > 0:
                    soma.setdefault((lj.marca, dia, int(hora)), []).append(pct)
    linhas = []
    for marca in dict.fromkeys(lj.marca for lj in lojas):
        for dia in DIAS:
            n_lojas = sum(1 for lj in lojas if lj.marca == marca
                          and any(lj.horarios_pico.get(dia, {}).values()))
            # Hora com poucas lojas abertas (ex.: 22h so no Goiania Shopping)
            # distorce a media da marca. Exige ao menos metade das lojas.
            minimo = max(1, -(-n_lojas // 2))
            curva = {h: sum(v) / len(v) for (m, d, h), v in soma.items()
                     if m == marca and d == dia and len(v) >= minimo}
            if not curva:
                continue
            ordem = sorted(curva, key=curva.get, reverse=True)
            linhas.append({
                "marca": marca, "dia": dia,
                "hora_pico": ordem[0], "movimento_pico_pct": round(curva[ordem[0]]),
                "horas_mais_cheias": ", ".join(f"{h:02d}h ({round(curva[h])}%)" for h in ordem[:top]),
                "hora_mais_vazia": min(curva, key=curva.get),
                "movimento_medio_pct": round(sum(curva.values()) / len(curva)),
                "lojas_com_dado": n_lojas,
                **{f"h{h:02d}": round(curva[h]) if h in curva else None for h in range(6, 24)},
            })
    return linhas


def exportar(lojas: list[Loja], nome: str) -> dict[str, Path]:
    SAIDA.mkdir(parents=True, exist_ok=True)
    arquivos: dict[str, Path] = {}
    cj = SAIDA / f"{nome}.json"
    cj.write_text(json.dumps([asdict(l) for l in lojas], ensure_ascii=False, indent=1), encoding="utf-8")
    arquivos["json"] = cj
    tab = tabelas(lojas)
    for k, linhas in tab.items():
        p = SAIDA / f"{nome}_{k}.csv"
        _escrever_csv(p, linhas)
        arquivos[k] = p
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        wb.remove(wb.active)
        for k, linhas in tab.items():
            ws = wb.create_sheet(k[:31])
            if not linhas:
                continue
            cab = list(linhas[0].keys())
            ws.append(cab)
            for l in linhas:
                ws.append([l[c] for c in cab])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, c in enumerate(cab, 1):
                larg = min(60, max(len(c), *(len(str(l[c] or "")) for l in linhas[:200])) + 2)
                ws.column_dimensions[get_column_letter(i)].width = larg
        px = SAIDA / f"{nome}.xlsx"
        wb.save(px)
        arquivos["xlsx"] = px
    except ImportError:
        pass
    return arquivos
